from openpyxl import Workbook
import csv
import json
import shutil
from py7zr import pack_7zarchive, unpack_7zarchive
from utils.settings import Settings
from enumeration.package_type import PackageType
from handler.file_handler import FileOsHandler
import cv2 as cv
from collections import defaultdict
from core.build_metadata import MetaDataCreator
from error.error_hub import ErrorHub
import gc

#Register the package to prevent register and threading problems for later on.
try:
    shutil.register_archive_format('7zip', pack_7zarchive, description='7zip archive')
    shutil.register_unpack_format('7zip', ['.7z'], unpack_7zarchive)
except Exception as e:
    pass

class Exporter():
    """Export class.
    Exports the accumulated data from the text handler and detection handler.
    Exports the data depending on the given preferences in the settings. (Excel etc. and compressed or not.)
    """
    def __init__(self):
        pass

        self.dir_path = Settings.retrieve_settings('directory_path')
        self.dir_name = Settings.retrieve_settings('temp_dir_name')

    def __save_val(self, val):
        """
        This function is used to filter out any symbols, that could cause
        the value inside of an excel file to create a calculation.
        """
        if isinstance(val, str) and val.startswith(('=', '+', '-', '@')):
            return f"'{val}"
        return val
    
    def __insert_metadata(self, filepaths: list, sheet):
        """
        This function simply inserts the metadata from the pdf file,
        into an excel sheet.
        """
        for file_name in filepaths:

            sheet.append(["title", "creator", "creationDate", "modDate",
                        "ISO216_A", "pdfWidth", "pdfHeight", "pdfWidth_mm", "pdfHeight_mm",
                        "pdfWidth_px", "pdfHeight_px", "scaled_height", "scaled_width", "file_name"])

            bundle = MetaDataCreator(file_name).get()._get()
            meta_values = []
            for v in bundle.values():
                meta_values.append(f"{self.__save_val(v)}")
            
            meta_values.append(f"{self.__save_val(file_name)}")
            sheet.append(meta_values)
            meta_values.clear()


    def export_to_excel(self, data_to_enter: list, origin_file_s: list) -> None:
        """
        This function exports the tesseract data, raw data, interpreted data and metadata into the excel file.

        :type data_to_enter: list
        :param data_to_enter: the data to enter into the excel file (amount depends on how many files are extracted).
        :type origin_file_s: list
        :param origin_file_s: list with the original file paths that were used for extraction.
        """
        wb = Workbook(write_only=True)
        sheet1 = wb.create_sheet("Tesseract")
        sheet2 = wb.create_sheet("Raw easyOCR data")
        sheet3 = wb.create_sheet("Interpreted")
        sheet4 = wb.create_sheet("Metadata")
        
        try:
            for data in data_to_enter:
                tesseract = data['tesseract']
                raw_data = data['raw_data']
                interpreted = data['interpreted']

                for i, file_name in enumerate(origin_file_s):
                    safe_file_name = f"{self.__save_val(file_name)}"

                    sheet1.append(["Type", "Text", "Left", "Top", "Width", "Height", "Bestand naam"])
                    for entry in tesseract[i]:
                        row = [entry.type, self.__save_val(entry.text), entry.x, entry.y, entry.w, entry.h, safe_file_name]
                        sheet1.append(row)

                    sheet2.append(["Position", "Text", "Confidence", "Symbol id", "Text offset Left", "Text offset Top", "Text offset Width", "Text offset Height",
                                        "Symbol offset Left", "Symbol offset Top", "Symbol offset Width", "Symbol offset Height", "Bestand naam"])
                    for entry in raw_data[i]:
                        x, y, w, h = cv.boundingRect(entry.offset)
                        x2, y2, w2, h2 = entry.symbol_offset
                        row = [self.__save_val(entry.grid_position), self.__save_val(entry.text), entry.conf, self.__save_val(entry.symbol_id),
                            x, y, w, h, x2, y2, w2, h2, safe_file_name]
                        sheet2.append(row)


                    sheet3.append(["Symbol id", "Letter code", "Tag number", "Left", "Top", "Width", "Height", "Bestand naam"])
                    for entry in interpreted[i]:
                        row = [self.__save_val(entry.symbol_id), self.__save_val(entry.text), self.__save_val(entry.tagnumber),
                            entry.x, entry.y, entry.w, entry.h, safe_file_name]
                        sheet3.append(row)
            
            self.__insert_metadata(origin_file_s, sheet4)
            wb.save(f"{self.dir_path}/{self.dir_name}/output.xlsx")
        except TypeError as tp:
            print(tp)
            ErrorHub.error_msg(f"Something went wrong while trying to export to excel. '{tp}'", "warning")
        finally:
            data_to_enter.clear()
            wb.close()
            gc.collect()

    def export_to_json(self, data_to_enter: list, origin_file_s: list) -> None:
        """This function exports the tesseract data, raw data, interpreted data and metadata into a json file.
        
        :type data_to_enter: list
        :param data_to_enter: the data to enter into the json file(s) (amount depends on how many files are extracted).
        :type origin_file_s: list
        :param origin_file_s: list with the original file paths that were used for extraction.
        """
        grouped_collection = defaultdict(list)
        for data in data_to_enter:
            tesseract = data['tesseract']
            raw_data = data['raw_data']
            interpreted = data['interpreted']

            for i, file_name in enumerate(origin_file_s):

                tess_boundary = []
                raw_boundary = []
                inter_boundary = []

                for entry in tesseract[i]:
                    tess_boundary.append({
                        "type": entry.type, 
                        "text": entry.text, 
                        "left": entry.x, 
                        "top": entry.y, 
                        "width": entry.w, 
                        "height": entry.h
                    })

                for entry in raw_data[i]:
                    x, y, w, h = cv.boundingRect(entry.offset)
                    x2, y2, w2, h2 = entry.symbol_offset
                    raw_boundary.append({
                        "position": entry.grid_position, 
                        "text": entry.text, 
                        "confidence": entry.conf, 
                        "symbol_id": entry.symbol_id, 
                        "text_offset_left": x, "text_offset_top": y, "text_offset_width": w, "text_offset_height": h, 
                        "symbol_offset_left": x2, "symbol_offset_top": y2, "symbol_offset_width": w2, "symbol_offset_height": h2
                    })
                    
                for entry in interpreted[i]:
                    inter_boundary.append({
                        "symbol_id": entry.symbol_id, 
                        "lettercode": entry.text, 
                        "tag_number": entry.tagnumber,
                        "left": entry.x, "top": entry.y, "width": entry.w, "height": entry.h
                    })
                    
                bundle = MetaDataCreator(file_name).get()._get() 

                grouped_collection[file_name].append({"tesseract": tess_boundary, "raw_easyocr_data": raw_boundary, "interpreted": inter_boundary, "metadata": [bundle]})

        _json = {"export": [{
            "collection": grouped_collection
        }]}

        file_path = f"{self.dir_path}/{self.dir_name}/output.json"

        try:
            with open(file_path, "w", encoding='utf-8') as file:
                json.dump(_json, file, ensure_ascii=False, indent=4)
        except FileNotFoundError:
            print("The directory to export too, does not exist. reverting to default, attempting once more.")
            ErrorHub.error_msg("The directory to export too, does not exist. reverting to default, attempting once more.", "information")
            FileOsHandler().create_dir(dir_path=self.dir_path, dir_name=self.dir_name)
            try:
                with open(file_path, "w", encoding='utf-8') as file:
                    json.dump(_json, file, ensure_ascii=False, indent=4)
            except json.JSONDecodeError as ex:
                print(ex)
                ErrorHub.error_msg(f"Something went wrong while trying to create a json file after the first attempt. '{ex}'", "critical")
        data_to_enter.clear()

    def export_to_csv(self, data_to_enter: list, origin_file_s: list) -> None:
        """This function exports the tesseract data, raw data, interpreted data and metadata into a csv file."""
        file_path = f"{self.dir_path}/{self.dir_name}/output.csv"
        
        filefieldname = ['file']
        fieldnames = ['type', 'text', 'left', 'top', 'width', 'height']
        fieldnames_raw = ['position', 'text', 'confidence', 'symbol_id', 'text_offset_left', 'text_offset_top', 'text_offset_width', 'text_offset_height',
                        'symbol_offset_left', 'symbol_offset_top', 'symbol_offset_width', 'symbol_offset_height']
        fieldnames_interpreted = ['symbol_id', 'lettercode', 'tag_number', 'left', 'top', 'width', 'height']
        fieldnames_metadata = ["title", "creator", "creationDate", "modDate", "ISO216_A", "pdfWidth", "pdfHeight", "pdfWidth_mm", "pdfHeight_mm",
                            "pdfWidth_px", "pdfHeight_px", "scaled_height", "scaled_width"]

        def write_csv_data():
            with open(file_path, "w", newline='', encoding='utf-8') as file:
                
                writer_file = csv.DictWriter(file, fieldnames=filefieldname)
                writer_tess = csv.DictWriter(file, fieldnames=fieldnames)
                writer_raw = csv.DictWriter(file, fieldnames=fieldnames_raw)
                writer_inter = csv.DictWriter(file, fieldnames=fieldnames_interpreted)
                writer_meta = csv.DictWriter(file, fieldnames=fieldnames_metadata)

                for data in data_to_enter:
                    tesseract = data['tesseract']
                    raw_data = data['raw_data']
                    interpreted = data['interpreted']

                    for i, file_name in enumerate(origin_file_s):
                        writer_file.writerow({"file": file_name})

                        for entry in tesseract[i]:
                            writer_tess.writerow({
                                "type": entry.type, "text": entry.text, "left": entry.x, 
                                "top": entry.y, "width": entry.w, "height": entry.h
                            })

                        for entry in raw_data[i]:
                            x, y, w, h = cv.boundingRect(entry.offset)
                            x2, y2, w2, h2 = entry.symbol_offset
                            writer_raw.writerow({
                                "position": entry.grid_position, "text": entry.text, "confidence": entry.conf, "symbol_id": entry.symbol_id, 
                                "text_offset_left": x, "text_offset_top": y, "text_offset_width": w, "text_offset_height": h, 
                                "symbol_offset_left": x2, "symbol_offset_top": y2, "symbol_offset_width": w2, "symbol_offset_height": h2
                            })
                            
                        for entry in interpreted[i]:
                            writer_inter.writerow({
                                "symbol_id": entry.symbol_id, "lettercode": entry.text, "tag_number": entry.tagnumber,
                                "left": entry.x, "top": entry.y, "width": entry.w, "height": entry.h
                            })

                        bundle = MetaDataCreator(file_name).get()._get()
                        writer_meta.writerow(bundle)

        try:
            write_csv_data()
        except FileNotFoundError:
            print("The directory to export to does not exist. Reverting to default, attempting once more.")
            ErrorHub.error_msg("The directory to export to does not exist. Reverting to default, try again.", "information")
            FileOsHandler().create_dir(dir_path=self.dir_path, dir_name=self.dir_name)
            try:
                write_csv_data()
            except Exception as ex:
                print(f"Second attempt failed: {ex}")
        data_to_enter.clear()
    
    def execute_export(self, data_to_enter: list, origin_file_s: list, package_type: PackageType = PackageType.DEFAULT) -> None:
        """This function is like a main hub for the export options (referring to csv, json and excel).
        If new type is added, it should be added here as well.
        Switches between types depending on the settings of the user.

        :type data_to_enter: list
        :param data_to_enter: the list of data that should be entered inside of the exported files
        :type origin_file_s: list
        :param origin_file_s: the list of original file names, where the data originates from.
        :type package_type: PackageType
        :param package_type: The PackageType referring to the export options from the user settings (is set to DEFAULT when empty)
        """
        FileOsHandler().create_dir(dir_path=self.dir_path, dir_name=self.dir_name)

        if len(Settings.retrieve_settings("chosen_export")) == 0:
            self.export_to_excel(data_to_enter=data_to_enter, origin_file_s=origin_file_s)
        
        match Settings.retrieve_settings("chosen_export"):
            case "excel":
                self.export_to_excel(data_to_enter=data_to_enter, origin_file_s=origin_file_s)
            case "json":
                self.export_to_json(data_to_enter=data_to_enter, origin_file_s=origin_file_s)
            case "csv":
                self.export_to_csv(data_to_enter=data_to_enter, origin_file_s=origin_file_s)

        self.execute_package_export(dir_path=self.dir_path, dir_name=self.dir_name, package_type=package_type)


    def execute_package_export(self, dir_path: str, dir_name: str, package_type: PackageType = PackageType.DEFAULT) -> None:
        """This function is the main hub for compressed export options, placing the exported files into the archives that are created here.
        It currently exports into 'zip', '7zip' and 'tar' but if other options are added, it should be added here as well.
        It creates an archive through the 'shutil' module, do keep in mind that the archive is flat. 
        (Meaning that there is no folder structure, just purely files.)

        dir_path and dir_name differentiate from another, the dir path is the full directory path where you want to export the compressed folder to.
        dir_name is the name of the eventual folder you will be archiving (exporting). So if dir_path is the path to Documents and dir_name 
        is for example called export_container (depending on user settings.). The files inside of the dir_name, will be archived inside of this folder.
        Documents/export_container -> Documents/(archived)export_container.

        Documents/ --(now archived)export_container/ --(flat, no folder structure inside the archive)files

        :type dir_path: string
        :param dir_path: the direct path to where the folder should be packaged.
        :type dir_name: string
        :param dir_name: the name of the folder that will be archived.
        :type package_type: PackageType
        :param package_type: the PackageType or the way of compression (DEFAULT if empty) (0 | ZIP, 1 | SEVENZ, 2 | TAR)
        """
        if isinstance(package_type, PackageType):
            package_type = package_type.value

        directory = f"{dir_path}/{dir_name}"

        match package_type:
            case 0:
                shutil.make_archive(base_name=directory, format="zip", root_dir=directory)
                FileOsHandler().remove_dir(directory)
            case 1:
                shutil.make_archive(base_name=directory, format="7zip", root_dir=directory)
                FileOsHandler().remove_dir(directory)
            case 2: 
                shutil.make_archive(base_name=directory, format="tar", root_dir=directory)
                FileOsHandler().remove_dir(directory)
            case 3:
                pass